"""Excel workbook export of the analysis."""

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER = PatternFill("solid", fgColor="1F3864")
SEV = {"High": "F8CBAD", "Medium": "FFE699", "Low": "DDEBF7"}


def to_excel(result: dict, ledger: pd.DataFrame, company: str, checklist: pd.DataFrame | None = None,
             extra: dict | None = None, summary: list | None = None) -> bytes:
    """extra: {sheet name: DataFrame} placed after the summary (case-level sheets for the credit note).
    summary: [(item, value)] rows added to the summary sheet."""
    buf = io.BytesIO()
    t = result["tables"]
    rows = [("Company", company), ("Prepared", datetime.now().strftime("%d %b %Y %H:%M")),
            ("Status", "DRAFT - AI-extracted figures, requires analyst review"), *(summary or []),
            ("Financial documents", len(ledger)),
            *[(f"{s} financial flags", int((result["flags"]["severity"] == s).sum()) if not result["flags"].empty else 0)
              for s in ("High", "Medium", "Low")]]
    sheets = {
        "Summary": pd.DataFrame(rows, columns=["Item", "Value"]),
        **(extra or {}),
        "Document checklist": checklist,
        "Financial KPIs":result["fin_kpis"].reset_index(drop=True) if not result["fin_kpis"].empty else None,
        "Banking KPIs": result["bank_kpis"],
        "GST KPIs": result["gst_kpis"],
        "Risk flags": result["flags"],
        "Financials (extracted)": t["financials"].reset_index() if not t["financials"].empty else None,
        "Bank months": t["bank"],
        "GST periods": t["gst"],
        "Director ITRs": t["itr"],
        "Source ledger": ledger,
    }
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None or df.empty:
                continue
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER
            for i, col in enumerate(df.columns, 1):
                width = max([len(str(col))] + [len(f"{v:,.2f}" if isinstance(v, float) else str(v)) for v in df[col].head(50)])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 70)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "#,##0.00" if abs(cell.value) < 1000 else "#,##0"
            if name in ("Risk flags", "All findings"):
                for r, sev in enumerate(df["severity"], 2):
                    for cell in ws[r]:
                        cell.fill = PatternFill("solid", fgColor=SEV[sev])
            ws.freeze_panes = "B2"
    return buf.getvalue()
